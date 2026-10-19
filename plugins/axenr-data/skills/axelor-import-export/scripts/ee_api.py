#!/usr/bin/env python3
"""Client minimal de l'import / export AOP EE (8.2.x).

Connexion à une instance HÉBERGÉE : aucun mot de passe ici. L'utilisateur ouvre une session dans son terminal
(curl, fichier de cookies ; recette dans references/import-ee-templates.md, « Se connecter »), puis exporte trois variables :
  AOS_URL         URL de l'instance, sans « / » final
  AOS_COOKIE_JAR  fichier de cookies (format Netscape écrit par curl) contenant JSESSIONID
  AOS_CSRF        valeur du cookie CSRF-TOKEN du fichier (à défaut, relue dans le fichier)
Dès que AOS_COOKIE_JAR est défini, le script n'appelle jamais /callback et ignore --cred et les variables
AXELOR_* ; AOS_URL est alors obligatoire.
  python3 ee_api.py fields Product > fields/Product.csv

Connexion à une instance LOCALE jetable : --cred <fichier json {"url", "username", "password"}>
(hors de tout dépôt, droits 600), ou variables AXELOR_URL, AXELOR_USER, AXELOR_PASSWORD. --cred se place
AVANT la sous-commande :
  ee_api.py --cred cred.json fields Product > fields/Product.csv

Lecture seule (aucune écriture en base) :
  info                             plateforme (EE ?), version AOP, modules, apps actives, plafond d'envoi,
                                   sièges de licence (plafond, comptes non bloqués, sièges libres)
  fields <Modele>                  relevé des champs de l'env cible (CSV) ; champs Studio marqués à part
  keys <Modele> [--domain D]       pour chaque fiche d'un référentiel : la clé à écrire dans un gabarit
  plan <Classeur.xlsx>             prévision ligne par ligne : CRÉATION ou MISE À JOUR de quelle fiche existante ;
                                   onglet User : comptes qui seront REBLOQUÉS faute de siège de licence
  links <Classeur.xlsx>            chaque valeur de lien existe-t-elle sur l'env (par sa clé) ou dans le classeur ?
  export <Modele> -f a,b.c -o F    export technique multi-onglets (fichier temporaire serveur)

Écrit une pièce jointe (MetaFile + fichier envoyé), rien d'autre :
  parse <Modele> <Classeur.xlsx>   envoie le fichier puis analyse à blanc les onglets et colonnes reconnus

Écrit les données :
  import <Modele> <Classeur.xlsx> --compute Onglet1,Onglet2 --confirm-write <URL> -o rapport.xlsx
    <Modele> = modèle racine ; tous les onglets du classeur dont le nom est un modèle sont importés, dans
    l'ordre de leurs dépendances. --confirm-write doit répéter l'URL de l'instance À L'IDENTIQUE (schéma,
    hôte, port, chemin) : pas d'import par mégarde sur une autre instance du même hôte.
"""
import argparse, csv, http.cookiejar, json, os, re, sys
import requests

def cookies_session_aos(chemin):
    """Cookies d'une session ouverte par curl (fichier au format Netscape), prêts pour un CookieJar.

    Lu à la main : curl note les cookies HttpOnly (JSESSIONID) « #HttpOnly_<hôte> » (les Python
    avant 3.10 les prennent pour des commentaires) et ceux de session « expires 0 » (MozillaCookieJar
    les croit expirés et n'envoie rien). Dans les deux cas la session serait vide, sans erreur."""
    cookies = []
    with open(chemin, encoding="utf-8") as f:
        for ligne in f:
            ligne = ligne.rstrip("\r\n")
            if ligne.startswith("#HttpOnly_"):
                ligne = ligne[len("#HttpOnly_"):]
            elif not ligne or ligne.startswith("#"):
                continue
            c = ligne.split("\t")
            if len(c) == 7:
                cookies.append(http.cookiejar.Cookie(
                    0, c[5], c[6], None, False, c[0], c[1] == "TRUE", c[0].startswith("."), c[2], True,
                    c[3] == "TRUE", None, True, None, None, {}))
    if not any(c.name == "JSESSIONID" for c in cookies):
        sys.exit("session : pas de JSESSIONID dans le fichier AOS_COOKIE_JAR (session absente ou expirée) : "
                 "l'utilisateur doit rouvrir la session (recette curl de references/import-ee-templates.md, « Se connecter »)")
    return cookies


class Client:
    def __init__(self, url, user=None, pwd=None):
        session = os.environ.get("AOS_COOKIE_JAR")  # session ouverte par curl : URL et identifiants ignorés
        if session:
            url = os.environ.get("AOS_URL")
            if not url:
                sys.exit("AOS_COOKIE_JAR est défini mais pas AOS_URL : exporter l'URL de l'instance (sans « / » final)")
        self.base = url.rstrip("/")
        self.s = requests.Session()
        if session:  # cookies et CSRF de la session, aucun /callback
            cookies = cookies_session_aos(session)
            for c in cookies:
                self.s.cookies.set_cookie(c)
            csrf = os.environ.get("AOS_CSRF") or next((c.value for c in cookies if c.name == "CSRF-TOKEN"), "")
            if not csrf:
                print("attention : ni AOS_CSRF ni cookie CSRF-TOKEN dans le fichier de cookies : les lectures "
                      "passeront, toute écriture sera refusée", file=sys.stderr)
            self.s.headers["X-CSRF-Token"] = csrf
            try:
                connecte = bool((self.s.get(self.base + "/ws/public/app/info").json().get("user") or {}).get("login"))
            except ValueError:
                connecte = False
            if not connecte:
                sys.exit("session refusée par le serveur (expirée ?) : l'utilisateur doit rouvrir la session "
                         "(recette curl de references/import-ee-templates.md, « Se connecter »)")
            return
        r = self.s.post(self.base + "/callback", json={"username": user, "password": pwd})
        if r.status_code != 200:
            sys.exit(f"connexion refusée ({r.status_code}) : identifiants, compte bloqué ou siège de licence")
        self.s.headers["X-CSRF-Token"] = self.s.cookies.get("CSRF-TOKEN") or ""

    def get(self, path, **kw):
        return self.s.get(self.base + path, **kw)

    def post(self, path, **kw):
        return self.s.post(self.base + path, **kw)

    def search(self, model, fields, domain=None, ctx=None, limit=500, offset=0):
        body = {"fields": fields, "limit": limit, "offset": offset}
        if domain:
            body["data"] = {"_domain": domain, "_domainContext": ctx or {}}
        r = self.post(f"/ws/rest/{fqn(self, model)}/search", json=body).json()
        if r.get("status") != 0:
            raise RuntimeError(f"search {model} : {json.dumps(r)[:400]}")
        return r.get("data") or [], r.get("total", 0)

    def search_all(self, model, fields, domain=None, ctx=None, page=500):
        rows, off = [], 0
        while True:
            part, total = self.search(model, fields, domain, ctx, limit=page, offset=off)
            rows += part; off += page
            if not part or off >= total:
                return rows

    def meta(self, model):
        r = self.get(f"/ws/meta/fields/{fqn(self, model)}").json()
        if r.get("status") != 0:
            raise RuntimeError(json.dumps(r)[:400])
        return r["data"]

    def upload(self, path):
        data = open(path, "rb").read()
        ftype = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                 if path.endswith(".xlsx") else "text/csv")
        r = self.post("/ws/files/upload", data=data, headers={
            "Content-Type": "application/octet-stream", "X-File-Name": os.path.basename(path),
            "X-File-Size": str(len(data)), "X-File-Offset": "0", "X-File-Type": ftype})
        meta = r.json()
        if r.status_code != 200 or "id" not in meta:
            raise RuntimeError(f"envoi refusé ({r.status_code}) : {json.dumps(meta)[:400]}")
        return meta["id"]

    def data_import(self, model, file_id, parse, compute):
        body = {"model": fqn(self, model), "file": {"id": file_id}, "parse": parse,
                "options": {"compute": {k: True for k in compute}, "flat": False}}
        r = self.post("/ws/data/import", json=body)
        try:
            return r.status_code, r.json()
        except ValueError:
            return r.status_code, {"raw": r.text[:600]}

    def download(self, file_name, out):
        r = self.get("/ws/data/download", params={"fileName": file_name})
        open(out, "wb").write(r.content)
        return out


_FQN = {}
def fqn(c, model):
    """Nom simple (« Product ») ou complet → nom complet, lu sur l'env cible (MetaModel)."""
    if "." in model:
        return model
    if model not in _FQN:
        r = c.post("/ws/rest/com.axelor.meta.db.MetaModel/search", json={
            "fields": ["fullName"], "data": {"_domain": "self.name = :n", "_domainContext": {"n": model}}}).json()
        rows = r.get("data") or []
        if len(rows) != 1:
            raise RuntimeError(f"modèle « {model} » introuvable ou ambigu sur l'env cible ({len(rows)} résultats)")
        _FQN[model] = rows[0]["fullName"]
    return _FQN[model]


def dasherize(name):
    """ProductFamily → product-family : préfixe des clés fabriquées par l'export (TypeUtils.makeImportId)."""
    return re.sub(r"(?<!^)(?=[A-Z])", "-", name).lower()


def unique_key(c, model, columns=None):
    """Champ par lequel l'import retrouve une fiche existante (miroir de TypeUtils.findUnique, AOP 8.2.3) :
    champ requis ET unique, de préférence code / email / emailAddress / …Seq, parmi les colonnes du fichier
    s'il est fourni ; à défaut importId."""
    props = [f for f in c.meta(model).get("fields") or [] if not f.get("target")]
    ru = [f["name"] for f in props if f.get("required") and f.get("unique")]
    cand = next((n for n in ru if n in ("code", "email", "emailAddress") or n.endswith("Seq")), None)
    cols = columns or []
    if cand and cand in cols:
        return cand
    for n in cols:
        if n in ru:
            return n
    if cand:
        return cand
    return ru[0] if ru else "importId"


def seats(c):
    """Sièges de licence Guardian : plafond lu dans AppLicense.licenses (« <nom> (<n>) »), comptes non bloqués."""
    try:
        lic, _ = c.search("com.axelor.enterprise.db.AppLicense", ["licenses", "appLicenseStatus", "expiryDate"], limit=5)
    except RuntimeError:
        lic = []
    total = None
    if lic:
        nums = [int(n) for x in lic for n in re.findall(r"\((\d+)\)", x.get("licenses") or "")]
        total = sum(nums) if nums else None
    _, active = c.search("com.axelor.auth.db.User", ["id"], domain="self.blocked = false or self.blocked is null", limit=1)
    return {"lignes": [x.get("licenses") for x in lic], "statut": [x.get("appLicenseStatus") for x in lic],
            "expiration": [x.get("expiryDate") for x in lic], "plafond": total, "comptes_non_bloques": active,
            "sieges_libres": None if total is None else total - active}


def connect(args):
    if os.environ.get("AOS_COOKIE_JAR"):  # session ouverte par curl : ni --cred ni mot de passe
        return Client(os.environ.get("AOS_URL"))
    if getattr(args, "cred", None):
        c = json.load(open(args.cred))
        return Client(c["url"], c["username"], c["password"])
    try:
        return Client(os.environ["AXELOR_URL"], os.environ["AXELOR_USER"], os.environ["AXELOR_PASSWORD"])
    except KeyError:
        sys.exit("instance locale : fournir --cred ou AXELOR_URL / AXELOR_USER / AXELOR_PASSWORD ; "
                 "instance hébergée : exporter AOS_URL, AOS_COOKIE_JAR et AOS_CSRF après avoir ouvert une session "
                 "par curl (recette de references/import-ee-templates.md, « Se connecter »)")


def cmd_info(c, a):
    info = c.get("/ws/public/app/info").json()
    app = info.get("application", {})
    mods = c.search_all("com.axelor.meta.db.MetaModule", ["name", "moduleVersion"])
    apps = c.search_all("com.axelor.studio.db.App", ["code", "active"])
    m = {x["name"]: x.get("moduleVersion") for x in mods}
    out = {"ee": "axelor-enterprise-edition" in m, "aopVersion": app.get("aopVersion"), "name": app.get("name"),
           "mode": app.get("mode"), "uploadMaxSizeMo": info.get("data", {}).get("upload", {}).get("maxSize"),
           "resetPasswordEnabled": app.get("resetPasswordEnabled"), "licence": seats(c), "modules": m,
           "appsActives": sorted(x["code"] for x in apps if x.get("active"))}
    print(json.dumps(out, ensure_ascii=False, indent=1))
    if not out["ee"]:
        print("ATTENTION : pas de module axelor-enterprise-edition : l'import EE n'existe pas ici", file=sys.stderr)


def cmd_fields(c, a):
    d = c.meta(a.model)
    w = csv.writer(sys.stdout, delimiter=";")
    w.writerow(["origine", "champ", "libelle", "type", "cible", "requis", "unique", "defaut", "selection"])
    for f in sorted(d.get("fields") or [], key=lambda x: x["name"]):
        sel = "|".join(str(s.get("value")) for s in (f.get("selectionList") or []))
        w.writerow(["conteneur studio" if f.get("type") == "JSON" else "natif", f["name"], f.get("title"),
                    f.get("type"), f.get("target") or "", "oui" if f.get("required") else "",
                    "oui" if f.get("unique") else "", f.get("defaultValue", ""), sel])
    for holder, jf in (d.get("jsonFields") or {}).items():
        for name, f in (jf or {}).items():
            w.writerow([f"studio ({holder})", name, f.get("title"), f.get("type"),
                        f.get("target") or f.get("jsonTarget") or "", "oui" if f.get("required") else "", "", "", ""])


def cmd_keys(c, a):
    key = unique_key(c, a.model)
    name_col = next((f["name"] for f in c.meta(a.model).get("fields") or [] if f.get("nameColumn")), "name")
    fields = ["importId"] + ([key] if key != "importId" else []) + [name_col]
    rows = c.search_all(a.model, fields, domain=a.domain)
    w = csv.writer(sys.stdout, delimiter=";")
    w.writerow(["id", "libelle", "importId", f"cle_unique ({key})", "colonne_de_lien", "valeur", "portee"])
    for x in rows:
        if key != "importId" and x.get(key):
            col, val, port = f"<rel>.{key}", x.get(key), "portable"
        elif x.get("importId"):
            col, val, port = "<rel>.importId", x["importId"], "portable si la même clé existe sur l'autre base"
        else:
            col, val, port = "<rel>.importId", f"{dasherize(a.model)}-{x['id']}", "LOCALE à cette base"
        w.writerow([x["id"], x.get(name_col), x.get("importId") or "", x.get(key) if key != "importId" else "",
                    col, val, port])
    print(f"# {len(rows)} fiches, clé unique : {key}", file=sys.stderr)


def cmd_plan(c, a):
    """Prévision sans écriture : pour chaque ligne des onglets principaux, la fiche que l'import retrouvera."""
    import openpyxl
    wb = openpyxl.load_workbook(a.file, read_only=True, data_only=True)
    w = csv.writer(sys.stdout, delimiter=";")
    w.writerow(["onglet", "ligne", "importId", "cle_unique", "valeur", "prevision", "fiche_existante"])
    counts, seat_rows = {}, []
    for ws in wb.worksheets:
        if "." in ws.title:
            continue
        try:
            fqn(c, ws.title)
        except RuntimeError:
            continue
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        hdr = [str(h) if h is not None else "" for h in rows[0]]
        key = unique_key(c, ws.title, hdr)
        for n, r in enumerate(rows[1:], start=2):
            d = dict(zip(hdr, r))
            if not any(v not in (None, "") for v in r):
                continue
            hit, how = None, ""
            if key != "importId" and d.get(key) not in (None, ""):
                f, _ = c.search(ws.title, ["importId", key] + (["blocked"] if ws.title == "User" else []),
                                domain=f"self.{key} = :v", ctx={"v": str(d[key])}, limit=2)
                if f:
                    hit, how = f[0], key
            if hit is None and d.get("importId") not in (None, ""):
                f, _ = c.search(ws.title, ["importId"] + (["blocked"] if ws.title == "User" else []),
                                domain="self.importId = :v", ctx={"v": str(d["importId"])}, limit=2)
                if f:
                    hit, how = f[0], "importId"
            if hit is None:
                prev = "CREATION"
            elif how != "importId" and hit.get("importId") != d.get("importId"):
                prev = f"ECRASE une fiche existante (par {how}) et réécrit sa clé d'import"
            else:
                prev = f"MISE A JOUR (par {how})"
            if ws.title == "User" and str(d.get("blocked")).strip().lower() == "false" \
                    and (hit is None or hit.get("blocked") is not False):
                seat_rows.append((ws.title, n, d.get("code")))
            label = prev.split(" (")[0]
            counts[label] = counts.get(label, 0) + 1
            w.writerow([ws.title, n, d.get("importId"), key, d.get(key) if key != "importId" else "", prev,
                        "" if hit is None else f"id {hit['id']} importId {hit.get('importId')}"])
    print(f"# {counts}", file=sys.stderr)
    if seat_rows:
        st = seats(c); free = st["sieges_libres"]
        print(f"# sièges : plafond {st['plafond']}, non bloqués {st['comptes_non_bloques']}, libres {free} ; "
              f"comptes à ouvrir par ce classeur : {len(seat_rows)}", file=sys.stderr)
        if free is not None and len(seat_rows) > free:
            for tab, n, code in seat_rows[free if free > 0 else 0:]:
                w.writerow([tab, n, "", "code", code, "SERA REBLOQUE (plus de siège de licence libre)", ""])
            print(f"# ATTENTION : {len(seat_rows) - max(free, 0)} compte(s) seront rebloqués à l'enregistrement, dans l'ordre du fichier", file=sys.stderr)


def cmd_links(c, a):
    """Chaque valeur de lien <rel>.<clé> des onglets principaux existe-t-elle sur l'env ou dans le classeur ?"""
    import openpyxl
    wb = openpyxl.load_workbook(a.file, read_only=True, data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            hdr = [str(h) if h is not None else "" for h in rows[0]]
            sheets[ws.title] = [dict(zip(hdr, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
    missing_total = 0
    for tab, data in sheets.items():
        if "." in tab:
            continue
        try:
            fields = {f["name"]: f for f in c.meta(tab).get("fields") or []}
        except RuntimeError:
            continue
        for col in (data[0].keys() if data else []):
            head, _, sub = col.partition(".")
            f = fields.get(head)
            if not sub or not f or not f.get("target"):
                continue
            tgt = f["target"].split(".")[-1]
            vals = {str(d[col]) for d in data if d.get(col) not in (None, "")}
            local = {str(d.get(sub)) for d in sheets.get(tgt, []) if d.get(sub) not in (None, "")}
            miss = []
            for v in sorted(vals - local):
                if sub == "importId" and re.fullmatch(re.escape(dasherize(tgt)) + r"-(\d+)", v):
                    dom, ctx = "self.id = :v", {"v": int(v.rsplit("-", 1)[1])}
                else:
                    dom, ctx = f"self.{sub} = :v", {"v": v}
                r, _ = c.search(tgt, ["id"], domain=dom, ctx=ctx, limit=1)
                if not r:
                    miss.append(v)
            missing_total += len(miss)
            print(f"{'OK      ' if not miss else 'MANQUANT'} {tab}.{col} → {tgt} : {len(vals)} valeur(s), "
                  f"{len(vals & local)} créée(s) par le classeur, {len(miss)} introuvable(s) {miss[:10]}")
    sys.exit(1 if missing_total else 0)


def cmd_export(c, a):
    body = {"model": fqn(c, a.model), "format": "xlsx",
            "options": {"fields": a.fields.split(","), "technical": True, "flat": False}}
    if a.domain:
        body["filter"] = {"_domain": a.domain}
    r = c.post("/ws/data/export", json=body).json()
    if "fileName" not in r:
        sys.exit(json.dumps(r)[:400])
    print(c.download(r["fileName"], a.out))


def cmd_parse(c, a):
    st, res = c.data_import(a.model, c.upload(a.file), True, [])
    if st != 200 or not isinstance(res, list):
        sys.exit(f"parse {st} : {json.dumps(res)[:600]}")
    for s in res:
        cols, resolved = s.get("columns") or [], s.get("resolved") or {}
        print(f"onglet {s['name']} → {s['model']}  ({len(cols)} colonnes)")
        for col in cols:
            print(f"   {'ok ' if col in resolved else 'NON'} {col}")
    print("ATTENTION : « ok » = chemin reconnu, pas valeur écrite : un lien par une clé non unique est « ok » et restera vide.")


def cmd_import(c, a):
    if a.confirm_write.rstrip("/") != c.base:
        sys.exit(f"écriture refusée : --confirm-write doit valoir exactement « {c.base} »")
    st, res = c.data_import(a.model, c.upload(a.file), False, [x for x in a.compute.split(",") if x])
    if st != 200 or "fileName" not in res:
        sys.exit(f"import {st} : {json.dumps(res)[:800]}")
    out = c.download(res["fileName"], a.out)
    import openpyxl
    wb = openpyxl.load_workbook(out, read_only=True)
    for ws in wb.worksheets[:2]:
        print(f"== {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 40:
                print("   …"); break
            print("  ", row)
    print(f"rapport : {out} — un rapport vert ne prouve rien : lancer verify_import.py")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--cred")
    sp = p.add_subparsers(dest="cmd", required=True)
    sp.add_parser("info")
    x = sp.add_parser("fields"); x.add_argument("model")
    x = sp.add_parser("keys"); x.add_argument("model"); x.add_argument("--domain")
    x = sp.add_parser("plan"); x.add_argument("file")
    x = sp.add_parser("links"); x.add_argument("file")
    x = sp.add_parser("export"); x.add_argument("model"); x.add_argument("-f", "--fields", required=True); x.add_argument("--domain"); x.add_argument("-o", "--out", required=True)
    x = sp.add_parser("parse"); x.add_argument("model"); x.add_argument("file")
    x = sp.add_parser("import"); x.add_argument("model"); x.add_argument("file"); x.add_argument("--compute", required=True, help="onglets à recalculer, ex. Product,User"); x.add_argument("--confirm-write", required=True, help="URL exacte de l'instance"); x.add_argument("-o", "--out", required=True)
    a = p.parse_args()
    c = connect(a)
    globals()["cmd_" + a.cmd](c, a)

if __name__ == "__main__":
    main()
