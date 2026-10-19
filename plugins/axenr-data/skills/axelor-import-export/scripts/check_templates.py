#!/usr/bin/env python3
"""Contrôle statique des classeurs de gabarits AVANT tout import (hors ligne, aucune connexion).

  check_templates.py CLASSEUR.xlsx […] [--fields-dir DIR] [--expected colonnes-attendues.csv] [--final]

Sans relevé (--fields-dir) : colonnes interdites (`id`, `<rel>.id`, mot de passe sur User), clés d'import
au motif fabriqué par l'export, doublons de clé, lignes d'exemple (--final), noms d'onglet > 31 caractères.
Avec --fields-dir (sorties de `ee_api.py fields <Modele>` nommées <Modele>.csv) : en plus, colonne inconnue
du modèle (ignorée en silence à l'import), champ Studio (décision projet), lien par un champ qui n'est ni
id, ni importId, ni la clé requise-unique de la cible (restera vide), booléen hors true/false ou vide,
onglet au nom de modèle sans relevé. --expected : chaque colonne attendue présente dans le bon onglet.
Sortie : BLOQUANT / ATTENTION / info. Code retour 1 s'il y a un bloquant.
"""
import argparse, csv, os, re, sys
import openpyxl

MODELE = re.compile(r"^[A-Z][A-Za-z0-9]*(\.[a-z][A-Za-z0-9]*)?$")   # « Product », « User.roles »
FABRIQUEE_GENERIQUE = re.compile(r"^[a-z]+(-[a-z0-9]+)*-\d+$")      # « product-92 », « product-family-12 »

def dasherize(name):
    return re.sub(r"(?<!^)(?=[A-Z])", "-", name).lower()

def load_fields(d):
    res = {}
    if d:
        for fn in os.listdir(d):
            if fn.endswith(".csv"):
                rows = list(csv.DictReader(open(os.path.join(d, fn), encoding="utf-8"), delimiter=";"))
                res[fn[:-4]] = {r["champ"]: r for r in rows}
    return res

def fabriquee(value, prefixes):
    v = str(value)
    if prefixes:
        return any(re.fullmatch(re.escape(p) + r"-\d+", v) for p in prefixes)
    return bool(FABRIQUEE_GENERIQUE.fullmatch(v))

def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("files", nargs="+"); p.add_argument("--expected"); p.add_argument("--fields-dir")
    p.add_argument("--final", action="store_true", help="classeur à importer : aucune ligne EXEMPLE, tout lien vérifié")
    a = p.parse_args()
    fields = load_fields(a.fields_dir)
    prefixes = {dasherize(m) for m in fields}
    bloq, warn, info, seen = [], [], [], {}
    for path in a.files:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        book = os.path.basename(path)
        # Nom d'onglet > 31 caractères : le moteur écrit un nom court et sa correspondance dans « __config__ »
        # (format name / value lu dans ExcelReader.loadSheetNames, 8.2.4 ; non rejoué)
        long_names = {}
        if "__config__" in wb.sheetnames:
            for r in list(wb["__config__"].iter_rows(values_only=True))[1:]:
                if r and r[0] and str(r[0]).startswith("sheet."):
                    long_names[str(r[0])[6:]] = str(r[1])
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows or ws.title == "__config__":
                continue
            title = long_names.get(ws.title, ws.title)
            where = f"{book} / {title}"
            hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
            seen[(book, title)] = set(hdr)
            if len(ws.title) > 31:
                bloq.append(f"{where} : nom d'onglet > 31 caractères (tronqué par Excel)")
            is_model = bool(MODELE.fullmatch(title))
            if not is_model:
                info.append(f"{where} : onglet non-modèle, ignoré à l'import (légende)")
                continue
            model, _, coll = title.partition(".")
            known = fields.get(model)
            if a.fields_dir and known is None:
                bloq.append(f"{where} : aucun relevé de « {model} » — lancer ee_api.py fields {model}")
            # --- contrôles sans relevé
            for col in hdr:
                if col == "id" or col.endswith(".id"):
                    bloq.append(f"{where} / {col} : id local d'une base — à retirer")
                if model == "User" and col in ("password", "transientPassword"):
                    bloq.append(f"{where} / {col} : mot de passe écrit en clair, compte inutilisable — à retirer")
            imp_cols = [i for i, h in enumerate(hdr) if h == "importId" or h.endswith(".importId")]
            ik = hdr.index("importId") if "importId" in hdr else None
            keys = {}
            for n, r in enumerate(rows[1:], start=2):
                if not any(v not in (None, "") for v in r):
                    continue
                for i in imp_cols:
                    if i < len(r) and r[i] not in (None, "") and fabriquee(r[i], prefixes):
                        (warn if hdr[i] != "importId" and not a.final else bloq).append(
                            f"{where} ligne {n} / {hdr[i]} : « {r[i]} » = clé fabriquée par l'export, id LOCAL"
                            + (" (acceptable uniquement dans le classeur généré pour CETTE base)" if hdr[i] != "importId" else ""))
                if ik is not None and ik < len(r) and r[ik] not in (None, ""):
                    k = str(r[ik])
                    if a.final and k.upper().startswith("EXEMPLE"):
                        bloq.append(f"{where} ligne {n} : ligne d'exemple restée dans le classeur final")
                    if not coll and k in keys:
                        bloq.append(f"{where} ligne {n} : clé « {k} » déjà vue ligne {keys[k]}")
                    keys.setdefault(k, n)
                elif not coll:
                    bloq.append(f"{where} ligne {n} : pas de clé d'import")
            if known is None:
                continue
            # --- contrôles avec relevé
            for col in hdr:
                if coll:
                    break
                head, _, sub = col.partition(".")
                f = known.get(head)
                if f is None:
                    if not (model == "User" and col in ("password", "transientPassword")):
                        bloq.append(f"{where} / {col} : champ inconnu du modèle sur l'env cible — ignoré en silence")
                    continue
                if f["origine"] != "natif":
                    bloq.append(f"{where} / {col} : champ Studio ou conteneur Studio — choix réservé au projet")
                if sub and f.get("cible"):
                    tgt = f["cible"].split(".")[-1]
                    tf = fields.get(tgt)
                    ok = sub in ("importId", "id") or (tf and tf.get(sub, {}).get("requis") == "oui"
                                                         and tf.get(sub, {}).get("unique") == "oui")
                    if tf is None:
                        (bloq if a.final else warn).append(f"{where} / {col} : relevé de {tgt} absent, clé de lien non vérifiée")
                    elif not ok:
                        bloq.append(f"{where} / {col} : lien par « {sub} », ni importId ni clé requise-unique de {tgt} — restera VIDE, rapport vert")
            if model == "User" and not coll:
                for need, why in (("blocked", "sans elle le compte arrive bloqué"), ("language", "compute ne pose pas la langue par défaut")):
                    if need not in hdr:
                        warn.append(f"{where} : pas de colonne {need} — {why}")
            decs = [i for i, h in enumerate(hdr) if known.get(h, {}).get("type") == "DECIMAL"]
            for n, r in enumerate(rows[1:], start=2):
                for i in decs:
                    v = r[i] if i < len(r) else None
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        if abs(v) >= 1000 and v != int(v):
                            bloq.append(f"{where} ligne {n} / {hdr[i]} : décimal {v} en cellule NUMÉRIQUE — ligne rejetée à l'import (écrire en texte avec un point)")
                        else:
                            warn.append(f"{where} ligne {n} / {hdr[i]} : décimal en cellule numérique — écrire en texte avec un point")
            bools = [i for i, h in enumerate(hdr) if known.get(h, {}).get("type") == "BOOLEAN"]
            for n, r in enumerate(rows[1:], start=2):
                if not any(v not in (None, "") for v in r):
                    continue
                for i in bools:
                    v = r[i] if i < len(r) else None
                    if v in (None, ""):
                        warn.append(f"{where} ligne {n} / {hdr[i]} : vide → NULL, la fiche échappe aux filtres « = faux »")
                    elif str(v).strip().lower() not in ("true", "false"):
                        bloq.append(f"{where} ligne {n} / {hdr[i]} : « {v} » vaudra FAUX (seul « true » est vrai)")
    if a.expected:
        for r in csv.reader(open(a.expected, encoding="utf-8"), delimiter=";"):
            if not r or r[0].startswith("#") or r[0] == "classeur":
                continue
            book, tab, col = r[0].strip(), r[1].strip(), r[2].strip()
            hit = [cols for (b, t), cols in seen.items() if t == tab and (not book or b.startswith(book))]
            if not hit:
                bloq.append(f"attendu {book} / {tab} : onglet absent")
            elif not any(col in cols for cols in hit):
                bloq.append(f"attendu {book} / {tab} / {col} : colonne absente")
    for m in info:
        print("info     ", m)
    for m in warn:
        print("ATTENTION", m)
    for m in bloq:
        print("BLOQUANT ", m)
    print(f"{len(bloq)} bloquant(s), {len(warn)} attention(s), {len(info)} info(s)")
    sys.exit(1 if bloq else 0)

if __name__ == "__main__":
    main()
